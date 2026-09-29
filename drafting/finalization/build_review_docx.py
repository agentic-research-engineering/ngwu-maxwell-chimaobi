from pathlib import Path
import re, zipfile, shutil, tempfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "Ngwu-Maxwell-Chimaobi-Thesis-Final-Review.docx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

chapters = [ROOT / "drafting" / f"chapter-0{i}.md" for i in range(1, 6)]
abstract_path = ROOT / "drafting/finalization/07-final-abstract.md"
bib_path = ROOT / "drafting/finalization/06-final-bibliography.md"

def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def inline_tokens(text):
    # Convert Markdown emphasis to Word run properties; URLs and punctuation
    # remain literal and no Markdown delimiter is emitted into the DOCX.
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            yield part[2:-2], False, True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            yield part[1:-1], True, False
        else:
            yield part, False, False

def add_inline(p, text, size=12):
    for chunk, italic, bold in inline_tokens(text):
        r = p.add_run(chunk)
        # Leave non-bold runs unset so Heading styles can supply boldness.
        set_font(r, size=size, italic=italic, bold=True if bold else None)

def configure_body_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.widow_control = True

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t); fld.append(r); paragraph._p.append(fld)

def set_pgnum(section, fmt, start):
    sectPr = section._sectPr
    el = sectPr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType"); sectPr.append(el)
    el.set(qn("w:fmt"), fmt); el.set(qn("w:start"), str(start))

def add_page(doc, heading, placeholders):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run(heading); set_font(r, size=14, bold=True)
    for text in placeholders:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(18)
        r = p.add_run(text); set_font(r, size=12, bold=True)

def add_toc_field(p):
    pPr = p._p.get_or_add_pPr()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt_run = OxmlElement("w:r"); txt = OxmlElement("w:t"); txt.text = "Table of contents will update when fields are refreshed."; txt_run.append(txt)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    p._p.extend([begin, instr, sep, txt_run, end])

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.5)
sec.different_first_page_header_footer = True
set_pgnum(sec, "lowerRoman", 1)
add_page_number(sec.footer.paragraphs[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"; normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
for name, size, align in [("Heading 1",14,WD_ALIGN_PARAGRAPH.CENTER),("Heading 2",12,WD_ALIGN_PARAGRAPH.LEFT),("Heading 3",12,WD_ALIGN_PARAGRAPH.LEFT)]:
    st=styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=None
    st._element.rPr.rFonts.set(qn("w:ascii"),"Times New Roman"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Times New Roman")
    st.paragraph_format.alignment=align; st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6); st.paragraph_format.keep_with_next=True

# Title page
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("A PHILOSOPHICAL INVESTIGATION INTO THE INSTITUTIONS OF STATE IN\nFRANCIS FUKUYAMA'S 'ORIGINS OF POLITICAL ORDER'"); set_font(r,size=14,bold=True)
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("BY"); set_font(r,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("MAXWELL CHIMAOBI NGWU"); set_font(r,bold=True)
for _ in range(3): doc.add_paragraph()
for line in ["DEPARTMENT OF PHILOSOPHY","BIGARD MEMORIAL SEMINARY, ENUGU","IN AFFILIATION WITH THE UNIVERSITY OF IBADAN","IBADAN, NIGERIA"]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(line); set_font(r,bold=True)
for _ in range(3): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("[PENDING HUMAN-SUPPLIED FINAL SUBMISSION MONTH AND YEAR]"); set_font(r,bold=True)

add_page(doc,"APPROVAL PAGE",[
    "[PENDING HUMAN-SUPPLIED APPROVAL PAGE CONTENT]",
    "[PENDING HUMAN-SUPPLIED APPROVAL WORDING]",
    "[PENDING HUMAN-SUPPLIED SIGNATORY NAME(S), TITLE(S), SIGNATURE LINE(S), AND DATE LINE(S)]"])
add_page(doc,"CERTIFICATION",[
    "[PENDING HUMAN-SUPPLIED CERTIFICATION CONTENT]",
    "[PENDING HUMAN-SUPPLIED CERTIFICATION WORDING]",
    "[PENDING HUMAN-SUPPLIED SIGNATORY NAME(S), TITLE(S), SIGNATURE LINE(S), AND DATE LINE(S)]"])
add_page(doc,"DEDICATION",["[PENDING HUMAN-SUPPLIED DEDICATION CONTENT]"])
add_page(doc,"ACKNOWLEDGEMENTS",["[PENDING HUMAN-SUPPLIED ACKNOWLEDGEMENTS CONTENT]"])

doc.add_page_break()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("TABLE OF CONTENTS"); set_font(r,size=14,bold=True)
p=doc.add_paragraph(); add_toc_field(p)

doc.add_page_break()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("ABSTRACT"); set_font(r,size=14,bold=True)
abstract = "\n".join(abstract_path.read_text().splitlines()[2:]).strip()
p=doc.add_paragraph(); configure_body_paragraph(p); add_inline(p,abstract)

# New main-text section
main = doc.add_section(WD_SECTION_START.NEW_PAGE)
main.page_width, main.page_height = Inches(8.5), Inches(11)
main.top_margin = main.bottom_margin = main.left_margin = main.right_margin = Inches(1)
main.header_distance = main.footer_distance = Inches(0.5)
main.footer.is_linked_to_previous = False
main.footer.paragraphs[0].clear(); add_page_number(main.footer.paragraphs[0])
set_pgnum(main, "decimal", 1)

footnotes = []
global_fn = 0

def footnote_ref(p, note_text):
    global global_fn
    global_fn += 1
    footnotes.append((global_fn,note_text))
    r = p.add_run()
    r._r.get_or_add_rPr().append(OxmlElement("w:rStyle"))
    r._r.rPr[-1].set(qn("w:val"), "FootnoteReference")
    ref = OxmlElement("w:footnoteReference"); ref.set(qn("w:id"), str(global_fn)); r._r.append(ref)

def add_body_with_notes(p,text,defs):
    pos=0
    for m in re.finditer(r'\[\^(\d+)\]',text):
        add_inline(p,text[pos:m.start()])
        footnote_ref(p,defs[m.group(1)])
        pos=m.end()
    add_inline(p,text[pos:])

for ci,path in enumerate(chapters):
    lines=path.read_text().splitlines()
    defs={}
    body=[]
    for line in lines:
        m=re.match(r'^\[\^(\d+)\]:\s*(.*)$',line)
        if m: defs[m.group(1)]=m.group(2)
        else: body.append(line)
    if ci>0: doc.add_page_break()
    i=0
    while i<len(body):
        line=body[i].strip()
        if not line: i+=1; continue
        if line.startswith('# '):
            title=line[2:].strip()
            # Combine chapter label and title for exact TOC entry while preserving restrained chapter-start styling.
            if title.startswith("CHAPTER ") and i+2<len(body) and body[i+1].strip()=="" and body[i+2].startswith('# '):
                title = title + ": " + body[i+2][2:].strip(); i+=2
            p=doc.add_paragraph(style="Heading 1"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p,title,size=14)
        elif line.startswith('## '):
            p=doc.add_paragraph(style="Heading 2"); add_inline(p,line[3:].strip())
        elif line.startswith('### '):
            p=doc.add_paragraph(style="Heading 3"); add_inline(p,line[4:].strip())
        else:
            p=doc.add_paragraph(); configure_body_paragraph(p); add_body_with_notes(p,line,defs)
        i+=1

# Bibliography
doc.add_page_break()
p=doc.add_paragraph(style="Heading 1"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_inline(p,"BIBLIOGRAPHY",size=14)
bib_lines=[x.strip() for x in bib_path.read_text().splitlines()[2:] if x.strip()]
for entry in bib_lines:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent=Inches(0.5); p.paragraph_format.first_line_indent=Inches(-0.5)
    p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(12)
    add_inline(p,entry)

# Ask Word/LibreOffice to update fields when opened.
settings=doc.settings._element
upd=OxmlElement("w:updateFields"); upd.set(qn("w:val"),"true"); settings.append(upd)
doc.save(OUT)

# Patch true footnotes into the package.
with tempfile.TemporaryDirectory() as td:
    td=Path(td)
    with zipfile.ZipFile(OUT) as z: z.extractall(td)
    nsmap={'w':W}
    root=etree.Element(qn('w:footnotes'),nsmap=nsmap)
    for fid,text in footnotes:
        fn=etree.SubElement(root,qn('w:footnote')); fn.set(qn('w:id'),str(fid))
        p=etree.SubElement(fn,qn('w:p'))
        pPr=etree.SubElement(p,qn('w:pPr'))
        sp=etree.SubElement(pPr,qn('w:spacing')); sp.set(qn('w:before'),'0'); sp.set(qn('w:after'),'0'); sp.set(qn('w:line'),'240'); sp.set(qn('w:lineRule'),'auto')
        r=etree.SubElement(p,qn('w:r')); rPr=etree.SubElement(r,qn('w:rPr'))
        etree.SubElement(rPr,qn('w:rStyle')).set(qn('w:val'),'FootnoteReference')
        etree.SubElement(r,qn('w:footnoteRef'))
        space=etree.SubElement(r,qn('w:t')); space.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); space.text=' '
        for chunk,italic,bold in inline_tokens(text):
            rr=etree.SubElement(p,qn('w:r')); rp=etree.SubElement(rr,qn('w:rPr'))
            fonts=etree.SubElement(rp,qn('w:rFonts')); fonts.set(qn('w:ascii'),'Times New Roman'); fonts.set(qn('w:hAnsi'),'Times New Roman')
            sz=etree.SubElement(rp,qn('w:sz')); sz.set(qn('w:val'),'20'); etree.SubElement(rp,qn('w:szCs')).set(qn('w:val'),'20')
            if italic: etree.SubElement(rp,qn('w:i'))
            if bold: etree.SubElement(rp,qn('w:b'))
            tt=etree.SubElement(rr,qn('w:t')); tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); tt.text=chunk
    (td/'word/footnotes.xml').write_bytes(etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes'))
    relp=td/'word/_rels/document.xml.rels'; relroot=etree.parse(str(relp)).getroot()
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    if not any(x.get('Type','').endswith('/footnotes') for x in relroot):
        ids=[int(x.get('Id')[3:]) for x in relroot if x.get('Id','').startswith('rId') and x.get('Id')[3:].isdigit()]
        rel=etree.SubElement(relroot,'{%s}Relationship'%relns); rel.set('Id',f'rId{max(ids)+1}'); rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes'); rel.set('Target','footnotes.xml')
        relp.write_bytes(etree.tostring(relroot,xml_declaration=True,encoding='UTF-8',standalone='yes'))
    ctp=td/'[Content_Types].xml'; ctroot=etree.parse(str(ctp)).getroot(); ctns='http://schemas.openxmlformats.org/package/2006/content-types'
    if not any(x.get('PartName')=='/word/footnotes.xml' for x in ctroot):
        ov=etree.SubElement(ctroot,'{%s}Override'%ctns); ov.set('PartName','/word/footnotes.xml'); ov.set('ContentType','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
        ctp.write_bytes(etree.tostring(ctroot,xml_declaration=True,encoding='UTF-8',standalone='yes'))
    tmpout=OUT.with_suffix('.tmp.docx')
    with zipfile.ZipFile(tmpout,'w',zipfile.ZIP_DEFLATED) as z:
        for f in td.rglob('*'):
            if f.is_file(): z.write(f,f.relative_to(td))
    tmpout.replace(OUT)

print(OUT)
print('footnotes',len(footnotes),'bibliography',len(bib_lines))
